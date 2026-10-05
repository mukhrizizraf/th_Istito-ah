"""Build the Excel data templates in assets/templates/ from TH.TEMPLATES.

Run from the repository root:  python tools/build_templates.py
Needs Node (to read th-data.js) and openpyxl.

Each workbook gets a Guide sheet (what it is, the column guide and how the
page uses the rows, in English and Bahasa Melayu), then one sheet per
template sheet with a styled header, sample rows, live formulas for the
worked-out columns and drop-down lists where a column has fixed values.
"""
import json
import os
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GREEN = PatternFill('solid', fgColor='1D6F42')
GREEN_2 = PatternFill('solid', fgColor='2B8A57')
CALC = PatternFill('solid', fgColor='F6E9C6')
WHITE_B = Font(bold=True, color='FFFFFF')
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=16, color='0A5A40')
MUTED = Font(italic=True, color='566960')
WRAP = Alignment(wrap_text=True, vertical='top')
NUMFMT = {'rm': '#,##0.00', 'int': '0', 'pct': '0"%"'}
TYPE = {'rm': 'RM', 'int': 'number', 'pct': '%', 'yn': 'Y / N', 'text': 'text'}


def formula(col, row, i):
    f = (row.get('_f') or {}).get(col['k']) or (col['f0'] if i == 0 and col['f0'] else None) or col['f']
    return '=' + f.replace('{r}', str(i + 2)).replace('{p}', str(i + 1)) if f else None


def guide_sheet(wb, tid, T):
    ws = wb.active
    ws.title = 'Guide'
    ws['A1'] = "Istito'ah Tracker: data template"
    ws['A1'].font = TITLE
    ws['A2'] = os.path.basename(T['file'])
    ws['A2'].font = MUTED
    r = 4
    for lang, label in (('en', 'What this is'), ('ms', 'Apakah ini')):
        ws.cell(r, 1, label).font = BOLD
        ws.cell(r, 2, T['intro'][lang]).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 48
        r += 1
    ws.cell(r, 1, 'Sample rows are made up. Keep the header row; replace the rows below it. '
                  'Gold cells are worked out by formula. / Baris contoh adalah rekaan. Kekalkan baris tajuk.').font = MUTED
    r += 2
    for c, h in enumerate(('Sheet', 'Column', 'Header', 'Type', 'Meaning (EN)', 'Maksud (BM)', 'Allowed values'), 1):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill = WHITE_B, GREEN
    r += 1
    for s in T['sheets']:
        for j, col in enumerate(s['cols']):
            vals = (s['name'], get_column_letter(j + 1), col['k'], TYPE.get(col['type'], col['type']) + (' (formula)' if col['f'] else ''),
                    col['d']['en'], col['d']['ms'], ', '.join(col['opts']) if col['opts'] else '')
            for c, v in enumerate(vals, 1):
                ws.cell(r, c, v).alignment = WRAP
            r += 1
    r += 1
    for c, h in enumerate(('Step', 'Sheet', 'Columns', 'What it counts (EN)', 'Apa yang dikira (BM)', 'Feeds', 'Menyalurkan'), 1):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill = WHITE_B, GREEN
    r += 1
    for k, st in enumerate(T['steps'], 1):
        vals = (str(k) + (' (proposed)' if st['proposed'] else ''), T['sheets'][st['sheet']]['name'], ', '.join(st['cols']),
                st['t']['en'] + ': ' + st['rule']['en'], st['t']['ms'] + ': ' + st['rule']['ms'], st['feeds']['en'], st['feeds']['ms'])
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v).alignment = WRAP
        r += 1
    if T['privacy']:
        r += 1
        ws.cell(r, 1, 'Privacy').font = BOLD
        ws.cell(r, 2, T['privacy']['en'] + '\n' + T['privacy']['ms']).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 60
    for c, w in zip('ABCDEFG', (22, 12, 26, 16, 54, 54, 34)):
        ws.column_dimensions[c].width = w


def data_sheet(wb, s):
    ws = wb.create_sheet(s['name'][:31])
    for j, col in enumerate(s['cols'], 1):
        cell = ws.cell(1, j, col['k'])
        cell.font, cell.fill = WHITE_B, (GREEN_2 if col['f'] else GREEN)
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(34, len(col['k']) + 4))
    for i, row in enumerate(s['rows']):
        for j, col in enumerate(s['cols'], 1):
            f = formula(col, row, i)
            v = row.get(col['k'])
            cell = ws.cell(i + 2, j, f if f else ('' if v is None else v))
            if f:
                cell.fill = CALC
            if col['type'] in NUMFMT and (f or isinstance(v, (int, float))):
                cell.number_format = NUMFMT[col['type']]
    ws.freeze_panes = 'A2'
    last = max(len(s['rows']) + 1, 200)
    for j, col in enumerate(s['cols'], 1):
        if col['opts'] and len(','.join(col['opts'])) < 250:
            dv = DataValidation(type='list', formula1='"' + ','.join(col['opts']) + '"', allow_blank=True)
            dv.error, dv.errorTitle = 'Pick a value from the list.', 'Not allowed'
            ws.add_data_validation(dv)
            letter = get_column_letter(j)
            dv.add(f'{letter}2:{letter}{last}')


def main():
    data = json.loads(subprocess.check_output(['node', os.path.join(ROOT, 'tools', 'export-templates.js')]))
    for tid, T in data.items():
        wb = Workbook()
        guide_sheet(wb, tid, T)
        for s in T['sheets']:
            data_sheet(wb, s)
        wb.active = 1
        out = os.path.join(ROOT, *T['file'].split('/'))
        os.makedirs(os.path.dirname(out), exist_ok=True)
        wb.save(out)
        print('wrote', T['file'])


if __name__ == '__main__':
    main()
